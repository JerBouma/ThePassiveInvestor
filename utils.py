import io
from lxml import html, etree
import pandas as pd
from openpyxl.chart import Reference, LineChart
from openpyxl.chart.axis import DateAxis
from openpyxl.styles import Alignment, Font
from openpyxl.drawing.image import Image as ExcelImage
import requests
import urllib3
import os
import sys


def resourcePath(relativePath):
    try:
        basePath = sys._MEIPASS
    except Exception:
        basePath = os.path.abspath(".")

    return os.path.join(basePath, relativePath)


def symbolCollector(input):
    if input[:4] == 'http':
        page = requests.get(input)
        tree = html.fromstring(page.content)
        table = tree.xpath('//table')
        symbol = pd.read_html(etree.tostring(table[0],method='html'))
        symbol = symbol[0]['Symbol'].to_list()
    else:
        symbol = pd.read_excel(input, header=None)[0].to_list()

    return symbol


def dataPlacer(data,sheet, startingRow, column, columnKey,
    columnValue, horizonalAlignmentKey=False, horizonalAlignmentValue=False,
    changeKeyDimensions=True, changeValueDimensions=True):

    maxLengthKey = 0
    maxLengthValue = 0

    for key, value in data.items():
        keyPosition = sheet.cell(column=column, row=startingRow, value=key)
        valuePosition = sheet.cell(column=column+1, row=startingRow, value=value)
        startingRow += 1

        if horizonalAlignmentKey:
            keyPosition.alignment = Alignment(horizontal=horizonalAlignmentKey) 
        
        if horizonalAlignmentValue:
            valuePosition.alignment = Alignment(horizontal=horizonalAlignmentValue)

        lengthKey = len(str(keyPosition.value))
        lengthValue = len(str(valuePosition.value))

        if lengthKey > maxLengthKey:
            maxLengthKey = lengthKey

        if lengthValue > maxLengthValue:
            maxLengthValue = lengthValue

    if changeKeyDimensions:
        sheet.column_dimensions[columnKey].width    = maxLengthKey * 1.2
    
    if changeValueDimensions:
        sheet.column_dimensions[columnValue].width  = maxLengthValue * 1.2


def imagePlacer(imageURL, sheet, location):
    try:
        http = urllib3.PoolManager()
        imageLocation = http.request('GET', imageURL)
        imageFile = io.BytesIO(imageLocation.data)
        image = ExcelImage(imageFile)
        sheet.add_image(image,location)
    except Exception:
        sheet[location] = "No image available"
        sheet[location].font = Font(italic=True)


def graphPlacer(ticker,stockSheet,stockData,
    sheet,minCol,minRow,maxCol):

    data = Reference(stockSheet, min_col=minCol+1, min_row=minRow, max_col=maxCol+1, max_row=len(stockData))
    cats = Reference(stockSheet, min_col=1, min_row=3, max_col=1, max_row=len(stockData))
            
    chart = LineChart()
    chart.title = None
    chart.legend = None
    chart.y_axis.title = "Stock Price"
    chart.y_axis.crossAx = 500
    chart.x_axis = DateAxis()
    chart.x_axis.number_format = 'yyyy'
    chart.x_axis.title = "Date"

    chart.add_data(data)
    chart.set_categories(cats)

    sheet.add_chart(chart, 'E4')


background = '#e8e8e8'
buttonColor = '#4a00a0'
buttonTextColour = 'white'
entryTextColour = 'gray'

defaultKeyStatisticsChoices = ['fundInceptionDate',
                               'category',
                               'totalAssets']

defaultsummaryDetailChoices = ['currency',
                               'navPrice',
                               'previousClose']

emptyRiskStatistics = {"year": 0,
                       "alpha": 0,
                       "beta": 0,
                       "meanAnnualReturn": 0,
                       "rSquared": 0,
                       "stdDev": 0,
                       "sharpeRatio": 0,
                       "treynorRatio": 0}
